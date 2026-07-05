"""Router `/efemerides` — endpoints REST de efemérides (feat-53.4).

Endpoints:

- `GET /efemerides` — lista todas, opcionalmente filtradas por tipo y
  relevancia.
- `GET /efemerides/proximas?dias=30&relevancia_minima=media` — las que
  caen en los próximos N días desde HOY. La fecha base es la del server
  (ART). Útil para dashboard.
- `GET /efemerides/{id}` — detalle.
- `POST /efemerides/{id}/generar-declaracion` — genera un proyecto de
  declaración usando el LLM configurado (con FakeLlmProvider si no hay
  Anthropic key). Tenant-scoped: el perfil del despacho del request
  influye en el tono.

Las efemérides son **públicas** (no tenant-scoped). La declaración
sí: usa el perfil opositor del despacho del request.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Annotated
from uuid import UUID

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi import APIRouter, HTTPException, Query, Response, status
from pydantic import BaseModel, Field

from praxis.api.deps import CurrentContext, LlmProviderDep, SessionDep
from praxis.application.use_cases.generar_declaracion_desde_efemeride import (
    EfemerideNoEncontrada,
    GenerarDeclaracionDesdeEfemeride,
)
from praxis.domain import (
    EFEMERIDE_TIPO_LABELS,
    Efemeride,
    RelevanciaEfemeride,
    TipoEfemeride,
)
from praxis.infrastructure.persistence.repositories import (
    SqlAlchemyEfemerideRepository,
    SqlAlchemyPerfilOpositorRepository,
)

router = APIRouter(prefix="/efemerides", tags=["efemerides"])


# ---------------------------------------------------------------------------
# Schemas
# ---------------------------------------------------------------------------


class EfemerideDTO(BaseModel):
    id: UUID
    mes: int
    dia: int
    fecha_corta: str
    titulo: str
    tipo: str
    tipo_label: str
    relevancia: str
    descripcion: str | None
    fuente: str | None
    areas_tematicas: list[str]
    anio_unico: int | None
    es_recurrente: bool

    @classmethod
    def from_domain(cls, ef: Efemeride) -> EfemerideDTO:
        return cls(
            id=ef.id,
            mes=ef.mes,
            dia=ef.dia,
            fecha_corta=ef.fecha_corta,
            titulo=ef.titulo,
            tipo=ef.tipo.value,
            tipo_label=EFEMERIDE_TIPO_LABELS[ef.tipo],
            relevancia=ef.relevancia.value,
            descripcion=ef.descripcion,
            fuente=ef.fuente,
            areas_tematicas=ef.areas_tematicas,
            anio_unico=ef.anio_unico,
            es_recurrente=ef.es_recurrente,
        )


class GenerarDeclaracionResponse(BaseModel):
    efemeride: EfemerideDTO
    tema_generado: str = Field(
        ..., description="Texto que se le pasó al LLM como contexto."
    )
    articulado: list[str]
    fundamentos: str
    modelo: str


class ExportarDocxRequest(BaseModel):
    titulo_efemeride: str
    articulado: list[str]
    fundamentos: str


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------


@router.get("", response_model=list[EfemerideDTO])
async def listar_efemerides(
    session: SessionDep,
    tipo: Annotated[str | None, Query(description="Filtro por tipo")] = None,
    relevancia: Annotated[str | None, Query(description="Filtro por relevancia")] = None,
) -> list[EfemerideDTO]:
    if tipo is not None and tipo not in {t.value for t in TipoEfemeride}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Tipo inválido: {tipo}",
        )
    if relevancia is not None and relevancia not in {
        r.value for r in RelevanciaEfemeride
    }:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Relevancia inválida: {relevancia}",
        )

    repo = SqlAlchemyEfemerideRepository(session)
    efemerides = await repo.listar_todas(tipo=tipo, relevancia=relevancia)
    return [EfemerideDTO.from_domain(ef) for ef in efemerides]


@router.get("/proximas", response_model=list[EfemerideDTO])
async def proximas_efemerides(
    session: SessionDep,
    dias: Annotated[int, Query(ge=1, le=180)] = 30,
    relevancia_minima: Annotated[str, Query()] = "media",
) -> list[EfemerideDTO]:
    if relevancia_minima not in {"alta", "media", "baja"}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"relevancia_minima inválida: {relevancia_minima}",
        )

    hoy = date.today()
    repo = SqlAlchemyEfemerideRepository(session)
    efemerides = await repo.proximas(
        desde_mes=hoy.month,
        desde_dia=hoy.day,
        dias=dias,
        relevancia_minima=relevancia_minima,
    )
    return [EfemerideDTO.from_domain(ef) for ef in efemerides]


@router.get("/{efemeride_id}", response_model=EfemerideDTO)
async def detalle_efemeride(
    efemeride_id: UUID,
    session: SessionDep,
) -> EfemerideDTO:
    repo = SqlAlchemyEfemerideRepository(session)
    ef = await repo.buscar_por_id(efemeride_id)
    if ef is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Efeméride no encontrada",
        )
    return EfemerideDTO.from_domain(ef)


@router.post(
    "/{efemeride_id}/generar-declaracion",
    response_model=GenerarDeclaracionResponse,
)
async def generar_declaracion(
    efemeride_id: UUID,
    session: SessionDep,
    llm: LlmProviderDep,
    ctx: CurrentContext,
) -> GenerarDeclaracionResponse:
    """Genera un proyecto de declaración a partir de la efeméride.

    Usa el `LlmProvider` configurado: AnthropicLlmProvider si hay API
    key, FakeLlmProvider (con plantillas keyword-based) si no.
    El perfil opositor del despacho del request influye en el tono.
    """
    efemerides = SqlAlchemyEfemerideRepository(session)
    perfiles = SqlAlchemyPerfilOpositorRepository(session)
    uc = GenerarDeclaracionDesdeEfemeride(
        efemerides=efemerides, perfiles=perfiles, llm=llm,
    )
    try:
        resultado = await uc.ejecutar(
            efemeride_id=efemeride_id,
            despacho_id=ctx.despacho.id,
        )
    except EfemerideNoEncontrada as exc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=str(exc),
        ) from exc

    return GenerarDeclaracionResponse(
        efemeride=EfemerideDTO.from_domain(resultado.efemeride),
        tema_generado=resultado.tema_generado,
        articulado=resultado.articulado,
        fundamentos=resultado.fundamentos,
        modelo=resultado.modelo,
    )


@router.post("/exportar-docx")
async def exportar_declaracion_docx(
    payload: ExportarDocxRequest,
    ctx: CurrentContext,
) -> Response:
    """Renderiza un proyecto de declaración como archivo Word editable.

    El frontend pasa el articulado + fundamentos ya generados (para no
    re-llamar al LLM) y el endpoint los compone en un .docx con
    encabezado del despacho.
    """
    doc = Document()

    # Estilo base
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # Encabezado: nombre del legislador titular del despacho
    legislador = ctx.despacho.legislador_titular_slug or "Despacho parlamentario"
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p_header.add_run(legislador.upper())
    run.bold = True

    doc.add_paragraph()

    # Título
    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_titulo.add_run("PROYECTO DE DECLARACIÓN")
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph()

    # Articulado
    for i, articulo in enumerate(payload.articulado, start=1):
        p = doc.add_paragraph()
        run = p.add_run(f"Artículo {i}°.- ")
        run.bold = True
        p.add_run(articulo)

    # Fórmula de cierre del articulado
    if payload.articulado:
        p_cierre = doc.add_paragraph()
        run = p_cierre.add_run(f"Artículo {len(payload.articulado) + 1}°.- ")
        run.bold = True
        p_cierre.add_run("Comuníquese al Poder Ejecutivo nacional.")

    doc.add_paragraph()

    # Fundamentos
    p_fund = doc.add_paragraph()
    p_fund.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_fund.add_run("FUNDAMENTOS")
    run.bold = True
    run.font.size = Pt(13)

    doc.add_paragraph()

    # Saludo + cuerpo de los fundamentos
    p_saludo = doc.add_paragraph()
    run = p_saludo.add_run("Señor Presidente:")
    run.bold = True

    for parrafo in payload.fundamentos.split("\n"):
        texto = parrafo.strip()
        if texto:
            doc.add_paragraph(texto)

    # Firma
    doc.add_paragraph()
    doc.add_paragraph()
    p_firma = doc.add_paragraph()
    p_firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_firma.add_run("_______________________________")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    nombre_archivo = (
        "proyecto_declaracion_"
        + payload.titulo_efemeride.lower()
            .replace(" ", "_")
            .replace("/", "_")[:60]
        + ".docx"
    )

    return Response(
        content=buf.read(),
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{nombre_archivo}"',
        },
    )
